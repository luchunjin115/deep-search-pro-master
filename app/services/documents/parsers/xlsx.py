"""Bounded XLSX extraction with Sheet, row, cell, and formula boundaries."""

from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta
from math import isfinite
from typing import TYPE_CHECKING, BinaryIO, Literal, Self, TypeAlias
from zipfile import ZipFile, ZipInfo

import openpyxl  # type: ignore[import-untyped]
from openpyxl.cell.cell import Cell  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.workbook.workbook import Workbook  # type: ignore[import-untyped]
from openpyxl.worksheet._read_only import (  # type: ignore[import-untyped]
    ReadOnlyWorksheet,
)
from pydantic import Field, model_validator

from app.schemas.common import M1Schema
from app.services.documents.parsers.base import (
    DocumentEncryptedError,
    DocumentLimitError,
    DocumentParseError,
    ParserWarning,
    SourceLocator,
)

if TYPE_CHECKING:
    from app.core.config import Settings

PARSER_NAME: Literal["openpyxl"] = "openpyxl"
PARSER_VERSION = f"m2-xlsx-v1+openpyxl-{openpyxl.__version__}"
_READ_CHUNK_BYTES = 1024 * 1024
_REQUIRED_MEMBERS = frozenset(
    {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"}
)

SpreadsheetScalar: TypeAlias = str | int | float | bool | None
CellDataType: TypeAlias = Literal[
    "empty",
    "text",
    "number",
    "boolean",
    "date",
    "formula",
    "error",
]


class XlsxCell(M1Schema):
    """One XLSX grid cell; formulas are preserved and never evaluated."""

    row_number: int = Field(ge=1, le=1_048_576)
    column_number: int = Field(ge=1, le=16_384)
    coordinate: str = Field(pattern=r"^[A-Z]{1,3}[1-9][0-9]{0,6}$", max_length=10)
    value: SpreadsheetScalar
    data_type: CellDataType
    formula: str | None = None
    cached_value: SpreadsheetScalar = None
    locator: SourceLocator

    @model_validator(mode="after")
    def validate_formula_and_locator(self) -> XlsxCell:
        if self.data_type == "formula":
            if self.formula is None or not self.formula.startswith("="):
                raise ValueError("formula cells require an Excel formula")
            if self.value != self.cached_value:
                raise ValueError("formula value must equal its cached value")
        elif self.formula is not None or self.cached_value is not None:
            raise ValueError("non-formula cells cannot contain formula metadata")
        if (
            self.locator.sheet_name is None
            or self.locator.row_number != self.row_number
            or self.locator.column_number != self.column_number
            or self.locator.cell_range != self.coordinate
        ):
            raise ValueError("XLSX cell locator does not match cell")
        return self


class XlsxRow(M1Schema):
    """One physical worksheet row, including rows whose cells are all empty."""

    row_number: int = Field(ge=1, le=1_048_576)
    is_empty: bool
    cells: list[XlsxCell] = Field(max_length=16_384)
    locator: SourceLocator

    @model_validator(mode="after")
    def validate_cells_and_locator(self) -> XlsxRow:
        if [cell.column_number for cell in self.cells] != list(
            range(1, len(self.cells) + 1)
        ) or any(cell.row_number != self.row_number for cell in self.cells):
            raise ValueError("XLSX cells must form a complete one-based row")
        if self.is_empty != all(_cell_is_empty(cell) for cell in self.cells):
            raise ValueError("XLSX row empty flag does not match cells")
        if (
            self.locator.sheet_name is None
            or self.locator.row_start != self.row_number
            or self.locator.row_end != self.row_number
        ):
            raise ValueError("XLSX row locator does not match row")
        return self


class XlsxSheet(M1Schema):
    """One worksheet in workbook order with a simple first-nonempty-row header."""

    sheet_number: int = Field(ge=1, le=1000)
    sheet_name: str = Field(min_length=1, max_length=31)
    state: Literal["visible", "hidden", "veryHidden"]
    row_count: int = Field(ge=0, le=1_048_576)
    column_count: int = Field(ge=0, le=16_384)
    non_empty_row_count: int = Field(ge=0, le=1_048_576)
    header_row_number: int | None = Field(default=None, ge=1, le=1_048_576)
    headers: list[str] = Field(max_length=16_384)
    rows: list[XlsxRow] = Field(max_length=1_048_576)
    locator: SourceLocator

    @model_validator(mode="after")
    def validate_shape_and_locator(self) -> XlsxSheet:
        if self.row_count != len(self.rows):
            raise ValueError("XLSX row_count does not match rows")
        if [row.row_number for row in self.rows] != list(
            range(1, self.row_count + 1)
        ):
            raise ValueError("XLSX rows must be a complete one-based sequence")
        if any(len(row.cells) != self.column_count for row in self.rows):
            raise ValueError("XLSX rows must match column_count")
        if self.non_empty_row_count != sum(not row.is_empty for row in self.rows):
            raise ValueError("XLSX non_empty_row_count does not match rows")
        expected_header = next(
            (row.row_number for row in self.rows if not row.is_empty),
            None,
        )
        if self.header_row_number != expected_header:
            raise ValueError("XLSX header row must be the first non-empty row")
        if self.headers and len(self.headers) != self.column_count:
            raise ValueError("XLSX headers must match column_count")
        if self.locator.sheet_name != self.sheet_name:
            raise ValueError("XLSX sheet locator does not match sheet")
        for row in self.rows:
            if row.locator.sheet_name != self.sheet_name:
                raise ValueError("XLSX row locator does not match sheet")
            for cell in row.cells:
                if cell.locator.sheet_name != self.sheet_name:
                    raise ValueError("XLSX cell locator does not match sheet")
        return self


class XlsxParseResult(M1Schema):
    """Safe XLSX parser output without paths, business-table writes, or evaluation."""

    parser_name: Literal["openpyxl"] = PARSER_NAME
    parser_version: str = Field(min_length=1, max_length=64)
    source_type: Literal["xlsx"] = "xlsx"
    sheet_count: int = Field(ge=1, le=1000)
    total_row_count: int = Field(ge=0)
    total_cell_count: int = Field(ge=0, le=2_000_000)
    formula_count: int = Field(ge=0, le=2_000_000)
    formulas_without_cached_value: int = Field(ge=0, le=2_000_000)
    character_count: int = Field(ge=0, le=20_000_000)
    sheets: list[XlsxSheet] = Field(min_length=1, max_length=1000)
    warnings: list[ParserWarning] = Field(max_length=1000)

    @model_validator(mode="after")
    def validate_counts(self) -> XlsxParseResult:
        if self.sheet_count != len(self.sheets):
            raise ValueError("XLSX sheet_count does not match sheets")
        if [sheet.sheet_number for sheet in self.sheets] != list(
            range(1, self.sheet_count + 1)
        ):
            raise ValueError("XLSX sheets must be a complete one-based sequence")
        rows = [row for sheet in self.sheets for row in sheet.rows]
        cells = [cell for row in rows for cell in row.cells]
        if self.total_row_count != len(rows) or self.total_cell_count != len(cells):
            raise ValueError("XLSX total counts do not match sheets")
        formulas = [cell for cell in cells if cell.formula is not None]
        if self.formula_count != len(formulas):
            raise ValueError("XLSX formula_count does not match cells")
        if self.formulas_without_cached_value != sum(
            cell.cached_value is None for cell in formulas
        ):
            raise ValueError("XLSX uncached formula count does not match cells")
        actual_characters = sum(_cell_character_count(cell) for cell in cells)
        if self.character_count != actual_characters:
            raise ValueError("XLSX character_count does not match cells")
        return self


class XlsxParser:
    """Read XLSX in streaming mode after bounded ZIP central-directory checks."""

    def __init__(
        self,
        *,
        max_source_bytes: int = 25 * 1024 * 1024,
        max_archive_members: int = 5000,
        max_uncompressed_bytes: int = 100 * 1024 * 1024,
        max_compression_ratio: int = 200,
        max_sheets: int = 100,
        max_rows_per_sheet: int = 100_000,
        max_columns: int = 500,
        max_cells: int = 500_000,
        max_extracted_characters: int = 5_000_000,
    ) -> None:
        limits = (
            max_source_bytes,
            max_archive_members,
            max_uncompressed_bytes,
            max_compression_ratio,
            max_sheets,
            max_rows_per_sheet,
            max_columns,
            max_cells,
            max_extracted_characters,
        )
        if any(limit <= 0 for limit in limits) or (
            max_archive_members > 20_000
            or max_uncompressed_bytes > 500 * 1024 * 1024
            or max_compression_ratio > 1000
            or max_sheets > 1000
            or max_rows_per_sheet > 1_048_576
            or max_columns > 16_384
            or max_cells > 2_000_000
            or max_extracted_characters > 20_000_000
        ):
            raise ValueError("XLSX parser limits must be positive and bounded")
        self._max_source_bytes = max_source_bytes
        self._max_archive_members = max_archive_members
        self._max_uncompressed_bytes = max_uncompressed_bytes
        self._max_compression_ratio = max_compression_ratio
        self._max_sheets = max_sheets
        self._max_rows_per_sheet = max_rows_per_sheet
        self._max_columns = max_columns
        self._max_cells = max_cells
        self._max_extracted_characters = max_extracted_characters

    @classmethod
    def from_settings(cls, settings: Settings) -> Self:
        return cls(
            max_source_bytes=settings.upload_max_file_size_bytes,
            max_archive_members=settings.xlsx_max_archive_members,
            max_uncompressed_bytes=settings.xlsx_max_uncompressed_bytes,
            max_compression_ratio=settings.xlsx_max_compression_ratio,
            max_sheets=settings.xlsx_max_sheets,
            max_rows_per_sheet=settings.xlsx_max_rows_per_sheet,
            max_columns=settings.xlsx_max_columns,
            max_cells=settings.xlsx_max_cells,
            max_extracted_characters=settings.xlsx_max_extracted_characters,
        )

    @property
    def parser_name(self) -> str:
        return PARSER_NAME

    @property
    def parser_version(self) -> str:
        return PARSER_VERSION

    def parse(self, stream: BinaryIO) -> XlsxParseResult:
        source = self._read_bounded(stream)
        self._validate_archive(source)
        formulas_workbook: Workbook | None = None
        values_workbook: Workbook | None = None
        try:
            formulas_workbook = openpyxl.load_workbook(
                io.BytesIO(source),
                read_only=True,
                data_only=False,
                keep_links=False,
            )
            values_workbook = openpyxl.load_workbook(
                io.BytesIO(source),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            return self._extract(formulas_workbook, values_workbook)
        except DocumentLimitError:
            raise
        except Exception:  # noqa: BLE001 - malformed OOXML details stay private.
            raise DocumentParseError from None
        finally:
            if formulas_workbook is not None:
                formulas_workbook.close()
            if values_workbook is not None:
                values_workbook.close()

    def _read_bounded(self, stream: BinaryIO) -> bytes:
        chunks: list[bytes] = []
        total = 0
        try:
            stream.seek(0)
            while True:
                remaining = self._max_source_bytes + 1 - total
                chunk = stream.read(min(_READ_CHUNK_BYTES, remaining))
                if chunk == b"":
                    break
                if not isinstance(chunk, (bytes, bytearray, memoryview)):
                    raise TypeError("binary stream required")
                binary = bytes(chunk)
                chunks.append(binary)
                total += len(binary)
                if total > self._max_source_bytes:
                    raise DocumentLimitError
        except DocumentLimitError:
            raise
        except Exception:  # noqa: BLE001 - arbitrary stream failures stay private.
            raise DocumentParseError from None
        return b"".join(chunks)

    def _validate_archive(self, source: bytes) -> None:
        try:
            with ZipFile(io.BytesIO(source)) as archive:
                members = archive.infolist()
                if len(members) > self._max_archive_members:
                    raise DocumentLimitError
                names: set[str] = set()
                total_uncompressed = 0
                for member in members:
                    self._validate_member(member, names)
                    total_uncompressed += member.file_size
                    if total_uncompressed > self._max_uncompressed_bytes:
                        raise DocumentLimitError
                if not _REQUIRED_MEMBERS.issubset(names):
                    raise DocumentParseError
        except (DocumentEncryptedError, DocumentLimitError, DocumentParseError):
            raise
        except Exception:  # noqa: BLE001 - ZIP implementation details stay private.
            raise DocumentParseError from None

    def _validate_member(self, member: ZipInfo, names: set[str]) -> None:
        name = member.filename
        parts = name.split("/")
        if (
            not name
            or "\\" in name
            or name.startswith("/")
            or any(part in {"", ".", ".."} for part in parts[:-1])
            or name in names
        ):
            raise DocumentParseError
        names.add(name)
        if member.flag_bits & 0x1:
            raise DocumentEncryptedError
        if member.file_size < 0 or member.compress_size < 0:
            raise DocumentParseError
        if member.file_size == 0:
            return
        if member.compress_size == 0:
            raise DocumentLimitError
        if member.file_size > member.compress_size * self._max_compression_ratio:
            raise DocumentLimitError

    def _extract(
        self,
        formulas_workbook: Workbook,
        values_workbook: Workbook,
    ) -> XlsxParseResult:
        if len(formulas_workbook.worksheets) > self._max_sheets:
            raise DocumentLimitError
        if formulas_workbook.sheetnames != values_workbook.sheetnames:
            raise DocumentParseError

        sheets: list[XlsxSheet] = []
        total_cells = 0
        formula_count = 0
        uncached_formula_count = 0
        character_count = 0
        warnings: list[ParserWarning] = []

        for sheet_number, (formula_sheet, value_sheet) in enumerate(
            zip(formulas_workbook.worksheets, values_workbook.worksheets),
            start=1,
        ):
            sheet, sheet_formulas, sheet_uncached, sheet_characters = (
                self._extract_sheet(
                    sheet_number,
                    formula_sheet,
                    value_sheet,
                    total_cells,
                )
            )
            total_cells += sheet.row_count * sheet.column_count
            formula_count += sheet_formulas
            uncached_formula_count += sheet_uncached
            character_count += sheet_characters
            if total_cells > self._max_cells or character_count > (
                self._max_extracted_characters
            ):
                raise DocumentLimitError
            sheets.append(sheet)
            if sheet.row_count == 0:
                warnings.append(
                    ParserWarning(
                        code="empty_sheet",
                        message="该Sheet没有可提取的单元格内容",
                        locator=SourceLocator(sheet_name=sheet.sheet_name),
                    )
                )

        return XlsxParseResult(
            parser_version=self.parser_version,
            sheet_count=len(sheets),
            total_row_count=sum(sheet.row_count for sheet in sheets),
            total_cell_count=total_cells,
            formula_count=formula_count,
            formulas_without_cached_value=uncached_formula_count,
            character_count=character_count,
            sheets=sheets,
            warnings=warnings,
        )

    def _extract_sheet(
        self,
        sheet_number: int,
        formula_sheet: ReadOnlyWorksheet,
        value_sheet: ReadOnlyWorksheet,
        prior_cells: int,
    ) -> tuple[XlsxSheet, int, int, int]:
        max_row = formula_sheet.max_row
        max_column = formula_sheet.max_column
        if max_row > self._max_rows_per_sheet or max_column > self._max_columns:
            raise DocumentLimitError
        if prior_cells + max_row * max_column > self._max_cells:
            raise DocumentLimitError

        rows: list[XlsxRow] = []
        formula_count = 0
        uncached_formula_count = 0
        character_count = 0
        formula_rows = formula_sheet.iter_rows(
            min_row=1,
            max_row=max_row,
            min_col=1,
            max_col=max_column,
        )
        value_rows = value_sheet.iter_rows(
            min_row=1,
            max_row=max_row,
            min_col=1,
            max_col=max_column,
        )

        for row_number, (formula_row, value_row) in enumerate(
            zip(formula_rows, value_rows),
            start=1,
        ):
            cells: list[XlsxCell] = []
            for column_number, (formula_cell, value_cell) in enumerate(
                zip(formula_row, value_row),
                start=1,
            ):
                coordinate = f"{get_column_letter(column_number)}{row_number}"
                formula = (
                    str(formula_cell.value)
                    if formula_cell.data_type == "f"
                    else None
                )
                cached_value = (
                    _normalize_scalar(value_cell.value) if formula is not None else None
                )
                value = (
                    cached_value
                    if formula is not None
                    else _normalize_scalar(formula_cell.value)
                )
                data_type = _data_type(formula_cell, value, formula)
                locator = SourceLocator(
                    sheet_name=formula_sheet.title,
                    row_number=row_number,
                    column_number=column_number,
                    cell_range=coordinate,
                )
                cell = XlsxCell(
                    row_number=row_number,
                    column_number=column_number,
                    coordinate=coordinate,
                    value=value,
                    data_type=data_type,
                    formula=formula,
                    cached_value=cached_value,
                    locator=locator,
                )
                cells.append(cell)
                character_count += _cell_character_count(cell)
                if character_count > self._max_extracted_characters:
                    raise DocumentLimitError
                if formula is not None:
                    formula_count += 1
                    if cached_value is None:
                        uncached_formula_count += 1

            row_end_column = get_column_letter(max_column)
            rows.append(
                XlsxRow(
                    row_number=row_number,
                    is_empty=all(_cell_is_empty(cell) for cell in cells),
                    cells=cells,
                    locator=SourceLocator(
                        sheet_name=formula_sheet.title,
                        cell_range=f"A{row_number}:{row_end_column}{row_number}",
                        row_start=row_number,
                        row_end=row_number,
                    ),
                )
            )

        if all(row.is_empty for row in rows):
            rows = []
            max_row = 0
            max_column = 0
            formula_count = 0
            uncached_formula_count = 0
            character_count = 0

        header_row = next((row for row in rows if not row.is_empty), None)
        headers = (
            [_cell_display_text(cell) for cell in header_row.cells]
            if header_row is not None
            else []
        )
        return (
            XlsxSheet(
                sheet_number=sheet_number,
                sheet_name=formula_sheet.title,
                state=formula_sheet.sheet_state,
                row_count=max_row,
                column_count=max_column,
                non_empty_row_count=sum(not row.is_empty for row in rows),
                header_row_number=(
                    header_row.row_number if header_row is not None else None
                ),
                headers=headers,
                rows=rows,
                locator=SourceLocator(sheet_name=formula_sheet.title),
            ),
            formula_count,
            uncached_formula_count,
            character_count,
        )


def _normalize_scalar(value: object) -> SpreadsheetScalar:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not isfinite(value):
            raise DocumentParseError
        return value
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    return str(value)


def _data_type(
    cell: Cell,
    value: SpreadsheetScalar,
    formula: str | None,
) -> CellDataType:
    if formula is not None:
        return "formula"
    if value is None:
        return "empty"
    if cell.data_type == "e":
        return "error"
    if cell.is_date:
        return "date"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    return "text"


def _cell_is_empty(cell: XlsxCell) -> bool:
    return cell.value is None and cell.formula is None


def _scalar_text(value: SpreadsheetScalar) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _cell_display_text(cell: XlsxCell) -> str:
    if cell.formula is not None:
        return _scalar_text(cell.cached_value) or cell.formula
    return _scalar_text(cell.value)


def _cell_character_count(cell: XlsxCell) -> int:
    values = (
        (cell.formula, _scalar_text(cell.cached_value))
        if cell.formula is not None
        else (_scalar_text(cell.value),)
    )
    return sum(
        not character.isspace()
        for value in values
        if value is not None
        for character in value
    )
