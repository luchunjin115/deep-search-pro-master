"""Synthetic XLSX and CSV inputs for M2 table-parser tests."""

from __future__ import annotations

import io
from datetime import date
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]


def make_structured_xlsx() -> bytes:
    """Create three Sheets with formulas, a blank row, a date, and an empty Sheet."""

    workbook = Workbook()
    quotes = workbook.active
    quotes.title = "合成报价"
    quotes.freeze_panes = "A2"
    quotes.append(["供应商", "SKU", "单价", "数量", "总价"])
    quotes.append(["合成供应商A", "SYNTH-001", 12.5, 10, "=C2*D2"])
    quotes.append([None, None, None, None, None])
    quotes.append(["合成供应商B", "SYNTH-002", 15, 5, "=C4*D4"])
    _style_sheet(quotes, (18, 16, 12, 10, 12))

    inspection = workbook.create_sheet("质检记录")
    inspection.freeze_panes = "A2"
    inspection.append(["日期", "状态", "备注"])
    inspection.append([date(2026, 8, 29), "通过", "全部内容均为合成演示数据"])
    inspection["A2"].number_format = "yyyy-mm-dd"
    _style_sheet(inspection, (14, 12, 30))

    empty = workbook.create_sheet("空Sheet")
    empty.sheet_state = "hidden"

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def make_empty_xlsx() -> bytes:
    output = io.BytesIO()
    workbook = Workbook()
    workbook.active.title = "空表"
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def make_large_dimension_xlsx(*, row: int, column: int) -> bytes:
    output = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(row=row, column=column, value="synthetic-limit")
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def make_xlsx_with_archive_payload(payload_size: int) -> bytes:
    output = io.BytesIO(make_empty_xlsx())
    with ZipFile(output, mode="a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("xl/synthetic-limit-payload.bin", b"A" * payload_size)
    return output.getvalue()


def make_xlsx_with_unsafe_member() -> bytes:
    output = io.BytesIO(make_empty_xlsx())
    with ZipFile(output, mode="a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("../synthetic-private.txt", b"not extracted")
    return output.getvalue()


def make_utf8_sig_semicolon_csv() -> bytes:
    text = (
        "供应商;备注;金额\r\n"
        '合成供应商A;"字段中包含;分号";12.5\r\n'
        ";;\r\n"
        "合成供应商B;仅用于演示;15\r\n"
    )
    return text.encode("utf-8-sig")


def make_gb18030_csv() -> bytes:
    text = "商品,市场,数量\r\n合成台灯,德国,125\r\n合成台灯,法国,80\r\n"
    return text.encode("gb18030")


def make_tab_csv() -> bytes:
    return b"sku\tstatus\nSYNTH-001\tactive\n"


def _style_sheet(sheet: object, widths: tuple[int, ...]) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
