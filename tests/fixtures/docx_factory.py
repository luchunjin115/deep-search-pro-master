"""Small synthetic DOCX files for deterministic parser tests and visual QA."""

from __future__ import annotations

import io
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from PIL import Image, ImageDraw

_TABLE_WIDTHS_DXA = (2700, 6660)


def make_structured_docx() -> bytes:
    """Create a compact synthetic guide with interleaved headings and tables."""

    document = Document()
    _apply_compact_reference_guide(document)

    document.add_heading("Synthetic Product Guide", level=1)
    document.add_paragraph(
        "Model SYNTH-LAMP-01 uses synthetic demonstration values only."
    )
    _add_table(
        document,
        (
            ("Field", "Value"),
            ("Voltage", "12 V"),
        ),
    )

    document.add_heading("Safety", level=2)
    document.add_paragraph("Disconnect power before synthetic maintenance.")
    _add_table(
        document,
        (
            ("Check", "Requirement"),
            ("Moisture", "Keep away from standing water"),
            ("Inspection", "Use synthetic checklist QS-01"),
        ),
    )

    document.add_paragraph("")
    document.add_heading("Cleaning", level=3)
    document.add_paragraph("Use a dry, lint-free cloth on the synthetic sample.")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_empty_docx() -> bytes:
    """Create a valid Word package with no extractable body text."""

    output = io.BytesIO()
    document = Document()
    _apply_compact_reference_guide(document)
    document.save(output)
    return output.getvalue()


def make_docx_with_archive_payload(payload_size: int) -> bytes:
    """Add a highly compressible unused member for archive-limit tests."""

    output = io.BytesIO(make_empty_docx())
    with ZipFile(output, mode="a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/synthetic-limit-payload.bin", b"A" * payload_size)
    return output.getvalue()


def make_docx_with_unsafe_member() -> bytes:
    """Add a traversal-like member that a safe parser must reject."""

    output = io.BytesIO(make_empty_docx())
    with ZipFile(output, mode="a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("../synthetic-private.txt", b"not extracted")
    return output.getvalue()


def make_docx_with_inline_images(*, repeated: bool = False) -> bytes:
    """Create header/body/footer content with text around inline image anchors."""

    document = Document()
    _apply_compact_reference_guide(document)
    document.sections[0].header.paragraphs[0].text = "CONTROL CODE: QC-VISUAL-17"
    document.sections[0].footer.paragraphs[0].text = "OWNER: QUALITY-LEAD"
    paragraph = document.add_paragraph()
    paragraph.add_run("BEFORE IMAGE")
    image_bytes = _make_test_card()
    paragraph.add_run().add_picture(io.BytesIO(image_bytes), width=Inches(2.5))
    if repeated:
        paragraph.add_run().add_picture(io.BytesIO(image_bytes), width=Inches(2.5))
    paragraph.add_run("AFTER IMAGE")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_docx_with_external_image_relationship() -> bytes:
    """Turn one embedded image relation into an external target for rejection tests."""

    source = make_docx_with_inline_images()
    input_stream = io.BytesIO(source)
    output = io.BytesIO()
    relationships_name = "word/_rels/document.xml.rels"
    namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    with (
        ZipFile(input_stream) as input_archive,
        ZipFile(
            output,
            mode="w",
            compression=ZIP_DEFLATED,
        ) as output_archive,
    ):
        for member in input_archive.infolist():
            payload = input_archive.read(member.filename)
            if member.filename == relationships_name:
                root = ElementTree.fromstring(payload)
                for relationship in root.findall(f"{{{namespace}}}Relationship"):
                    if relationship.get("Type", "").endswith("/image"):
                        relationship.set("Target", "https://invalid.example/image.png")
                        relationship.set("TargetMode", "External")
                payload = ElementTree.tostring(
                    root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            output_archive.writestr(member, payload)
    return output.getvalue()


def make_docx_with_corrupt_image() -> bytes:
    """Keep the image relationship but replace its package bytes with invalid data."""

    source = make_docx_with_inline_images()
    input_stream = io.BytesIO(source)
    output = io.BytesIO()
    with (
        ZipFile(input_stream) as input_archive,
        ZipFile(output, mode="w", compression=ZIP_DEFLATED) as output_archive,
    ):
        for member in input_archive.infolist():
            payload = input_archive.read(member.filename)
            if member.filename.startswith("word/media/"):
                payload = b"not-a-valid-image"
            output_archive.writestr(member, payload)
    return output.getvalue()


def _make_test_card() -> bytes:
    image = Image.new("RGB", (640, 180), "white")
    drawing = ImageDraw.Draw(image)
    drawing.text((20, 60), "ACTION: QUARANTINE 12 PCS", fill="black")
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _apply_compact_reference_guide(document: Document) -> None:
    """Apply the skill's compact_reference_guide tokens explicitly."""

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 14, 7),
        "Heading 3": (12, "1F4D78", 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def _add_table(document: Document, values: tuple[tuple[str, str], ...]) -> None:
    table = document.add_table(rows=len(values), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_table_geometry(table)

    for row_number, values_row in enumerate(values):
        for column_number, value in enumerate(values_row):
            cell = table.cell(row_number, column_number)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if column_number == 0
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.25
            if row_number == 0:
                for run in paragraph.runs:
                    run.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EEF5")
                cell._tc.get_or_add_tcPr().append(shading)


def _set_table_geometry(table: object) -> None:
    table_xml = table._tbl
    table_properties = table_xml.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.insert(0, table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), "9360")

    table_indent = table_properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:w"), "120")
    table_indent.set(qn("w:type"), "dxa")

    for grid_column, width in zip(table_xml.tblGrid.gridCol_lst, _TABLE_WIDTHS_DXA):
        grid_column.set(qn("w:w"), str(width))

    for row in table.rows:
        for cell, width in zip(row.cells, _TABLE_WIDTHS_DXA):
            cell.width = Twips(width)
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.get_or_add_tcW()
            cell_width.type = "dxa"
            cell_width.w = width
            margins = cell_properties.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell_properties.append(margins)
            for side, value in (
                ("top", 80),
                ("bottom", 80),
                ("start", 120),
                ("end", 120),
            ):
                element = margins.find(qn(f"w:{side}"))
                if element is None:
                    element = OxmlElement(f"w:{side}")
                    margins.append(element)
                element.set(qn("w:w"), str(value))
                element.set(qn("w:type"), "dxa")
