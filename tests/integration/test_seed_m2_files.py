from __future__ import annotations

import io
import json
from collections.abc import Generator
from pathlib import Path
from zipfile import ZipFile

import pytest
from alembic import command
from alembic.config import Config
from docx import Document as WordDocument
from docx.shared import Inches
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import Engine, delete, func, select
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.db.session import create_database_runtime
from app.models.catalog import ProductVariant
from app.models.identity import Role, Tenant
from app.models.inventory import InventorySnapshot, Warehouse
from app.models.knowledge import Document, DocumentAcl, DocumentVersion, StoredFile
from app.services.documents.parsers import CsvParser, DocxParser, PdfParser, XlsxParser
from app.services.documents.parsers.base import SourceLocator
from app.services.storage import LocalStorageBackend
from scripts.seed_m1 import deterministic_id as m1_deterministic_id
from scripts.seed_m2_files import (
    DEFAULT_SEED_PATH,
    M2SeedDataMismatchError,
    build_manifest,
    generate_sources,
    load_seed_definition,
    seed_m2_files,
)

_FIXED_ZIP_TIME = (2020, 1, 1, 0, 0, 0)


@pytest.fixture
def postgres_settings(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Settings:
    settings = Settings(  # type: ignore[call-arg]
        _env_file=".env.example",
        app_env="test",
        local_storage_root=tmp_path / "storage",
    )
    monkeypatch.setenv("DATABASE_URL", settings.database_url)
    return settings


@pytest.fixture
def postgres_engine(postgres_settings: Settings) -> Generator[Engine, None, None]:
    runtime = create_database_runtime(postgres_settings)
    yield runtime.engine
    runtime.engine.dispose()


def _clean_seed_rows(engine: Engine, data: dict[str, object]) -> None:
    tenant_id = m1_deterministic_id(
        "m1-v1",
        "tenant",
        str(data["tenant_name"]),
    )
    role_ids = [
        m1_deterministic_id("m1-v1", "role", role_name)
        for role_name in ("company_owner", "product_scout", "amazon_operator")
    ]
    with engine.begin() as connection:
        connection.execute(delete(Document).where(Document.tenant_id == tenant_id))
        connection.execute(delete(StoredFile).where(StoredFile.tenant_id == tenant_id))
        connection.execute(delete(Tenant).where(Tenant.id == tenant_id))
        connection.execute(delete(Role).where(Role.id.in_(role_ids)))


def test_m2_source_generation_is_deterministic_and_locators_match_parsers() -> None:
    data = load_seed_definition(DEFAULT_SEED_PATH)
    first = generate_sources(data)
    second = generate_sources(data)
    manifest = build_manifest(data, first)

    assert [source.content for source in first] == [source.content for source in second]
    assert [source.sha256 for source in first] == [source.sha256 for source in second]
    assert manifest["seeded_row_counts"] == {
        "files": 5,
        "documents": 5,
        "document_versions": 5,
        "document_acl": 5,
    }
    assert [source.definition["format"] for source in first] == [
        "pdf",
        "docx",
        "docx",
        "xlsx",
        "csv",
    ]
    for source in first:
        assert b"API_KEY" not in source.content
        assert b"PRIVATE KEY" not in source.content
        if source.definition["format"] in {"docx", "xlsx"}:
            with ZipFile(io.BytesIO(source.content)) as archive:
                assert all(member.date_time == _FIXED_ZIP_TIME for member in archive.infolist())
        for fact in source.definition["golden_facts"]:
            SourceLocator.model_validate(fact["locator"])

    sources = {source.definition["key"]: source for source in first}
    manual = PdfParser().parse(io.BytesIO(sources["mushroom_lamp_manual"].content))
    assert manual.page_count == 2
    assert "额定电压：220 V" in manual.pages[0].text
    assert "清洁前拔下插头" in manual.pages[1].text

    sop = DocxParser().parse(io.BytesIO(sources["quality_inspection_sop"].content))
    assert sop.paragraph_count == 16
    assert sop.table_count == 1
    assert sop.blocks[9].locator.paragraph_number == 10
    assert sop.blocks[9].locator.heading_path == ["2. 抽样规则"]
    sop_table = next(block for block in sop.blocks if block.kind == "table")
    assert sop_table.rows[1].cells[2].text == "立即隔离整批并升级质量负责人"
    assert sop_table.rows[1].cells[2].locator.row_number == 2
    assert sop_table.rows[1].cells[2].locator.column_number == 3

    compliance = DocxParser().parse(
        io.BytesIO(sources["de_compliance_checklist"].content)
    )
    compliance_disclaimer = compliance.blocks[6]
    assert compliance_disclaimer.kind == "paragraph"
    assert "演示资料，不构成法律或认证意见" in compliance_disclaimer.text
    compliance_table = next(
        block for block in compliance.blocks if block.kind == "table"
    )
    assert compliance_table.rows[3].cells[2].text == "缺少正式证明"

    quotes = XlsxParser().parse(io.BytesIO(sources["supplier_quotes"].content))
    assert [sheet.sheet_name for sheet in quotes.sheets] == [
        "供应商报价",
        "报价说明",
    ]
    assert [sheet.row_count for sheet in quotes.sheets] == [4, 5]
    assert quotes.sheets[0].rows[2].cells[5].value == 20.9
    assert quotes.sheets[0].rows[1].cells[7].formula == "=F2+G2"
    assert quotes.sheets[0].rows[1].cells[7].cached_value is None

    for key in ("quality_inspection_sop", "de_compliance_checklist"):
        word = WordDocument(io.BytesIO(sources[key].content))
        section = word.sections[0]
        assert section.page_width == Inches(8.5)
        assert section.page_height == Inches(11)
        assert section.top_margin == Inches(1)
        assert section.right_margin == Inches(1)
        assert section.bottom_margin == Inches(1)
        assert section.left_margin == Inches(1)
        assert word.tables
        assert all(table.autofit is False for table in word.tables)

    workbook = load_workbook(io.BytesIO(sources["supplier_quotes"].content), data_only=False)
    try:
        assert workbook.sheetnames == ["供应商报价", "报价说明"]
        for sheet in workbook.worksheets:
            assert sheet.freeze_panes == "A2"
            assert sheet.sheet_view.showGridLines is False
            assert sheet.max_row > 1
            assert sheet.max_column > 1
        assert workbook["供应商报价"]["H2"].value == "=F2+G2"
        assert workbook["供应商报价"].column_dimensions["A"].width == 18
    finally:
        workbook.close()

    operations = CsvParser().parse(io.BytesIO(sources["monthly_operations"].content))
    assert operations.encoding == "utf-8-sig"
    assert operations.delimiter == ","
    assert operations.row_count == 15
    assert operations.rows[6].values == [
        "2026-06",
        "DE",
        "LR-TL-MUSH-OR01",
        "139",
        "710",
        "7",
    ]
    assert operations.rows[7].values[1] == "FR"
    assert operations.rows[12].values[1] == "FR"


def test_seed_m2_files_is_repeatable_in_storage_and_postgresql(
    postgres_settings: Settings,
    postgres_engine: Engine,
    tmp_path: Path,
) -> None:
    command.upgrade(Config("alembic.ini"), "head")
    data = load_seed_definition(DEFAULT_SEED_PATH)
    manifest_path = tmp_path / "m2_manifest.json"
    m1_manifest_path = tmp_path / "m1_manifest.json"
    _clean_seed_rows(postgres_engine, data)

    try:
        first = seed_m2_files(
            postgres_settings,
            manifest_path=manifest_path,
            m1_manifest_path=m1_manifest_path,
        )
        first_manifest_bytes = manifest_path.read_bytes()
        second = seed_m2_files(
            postgres_settings,
            manifest_path=manifest_path,
            m1_manifest_path=m1_manifest_path,
        )
        second_manifest_bytes = manifest_path.read_bytes()

        assert first.manifest == second.manifest
        assert first_manifest_bytes == second_manifest_bytes
        assert json.loads(second_manifest_bytes) == second.manifest
        assert first.manifest["data_classification"] == "synthetic_demo_data"
        assert "不构成法律" in first.manifest["disclaimer"]

        tenant_id = m1_deterministic_id(
            "m1-v1",
            "tenant",
            data["tenant_name"],
        )
        with Session(postgres_engine) as session:
            counts = {
                "files": session.scalar(
                    select(func.count())
                    .select_from(StoredFile)
                    .where(StoredFile.tenant_id == tenant_id)
                ),
                "documents": session.scalar(
                    select(func.count())
                    .select_from(Document)
                    .where(Document.tenant_id == tenant_id)
                ),
                "document_versions": session.scalar(
                    select(func.count())
                    .select_from(DocumentVersion)
                    .where(DocumentVersion.tenant_id == tenant_id)
                ),
                "document_acl": session.scalar(
                    select(func.count())
                    .select_from(DocumentAcl)
                    .where(DocumentAcl.tenant_id == tenant_id)
                ),
            }
            files = list(
                session.scalars(
                    select(StoredFile)
                    .where(StoredFile.tenant_id == tenant_id)
                    .order_by(StoredFile.original_name)
                )
            )
            documents = list(
                session.scalars(
                    select(Document)
                    .where(Document.tenant_id == tenant_id)
                    .order_by(Document.title)
                )
            )
            versions = list(
                session.scalars(
                    select(DocumentVersion).where(
                        DocumentVersion.tenant_id == tenant_id
                    )
                )
            )
            available = session.scalar(
                select(
                    InventorySnapshot.on_hand
                    - InventorySnapshot.reserved
                    - InventorySnapshot.unsellable
                )
                .join(
                    ProductVariant,
                    ProductVariant.id == InventorySnapshot.variant_id,
                )
                .join(Warehouse, Warehouse.id == InventorySnapshot.warehouse_id)
                .where(
                    ProductVariant.sku == "LR-TL-MUSH-OR01",
                    Warehouse.code == "DE-FRA",
                )
                .order_by(InventorySnapshot.snapshot_at.desc())
                .limit(1)
            )

        assert counts == first.manifest["seeded_row_counts"]
        assert all(row.status == "uploaded" for row in files)
        assert all(row.deleted_at is None for row in files)
        assert all(row.active_version_id is None for row in documents)
        assert all(row.parse_status == "pending" for row in versions)
        assert all(row.index_status == "pending" for row in versions)
        assert available == 125

        storage = LocalStorageBackend(postgres_settings.local_storage_root)
        generated = generate_sources(data)
        for source in generated:
            with storage.open(source.storage_key) as stream:
                assert stream.read() == source.content
    finally:
        _clean_seed_rows(postgres_engine, data)


def test_seed_rejects_a_conflicting_existing_storage_object(
    postgres_settings: Settings,
    postgres_engine: Engine,
    tmp_path: Path,
) -> None:
    command.upgrade(Config("alembic.ini"), "head")
    data = load_seed_definition(DEFAULT_SEED_PATH)
    source = generate_sources(data)[0]
    storage = LocalStorageBackend(postgres_settings.local_storage_root)
    storage.put(source.storage_key, io.BytesIO(b"conflicting synthetic bytes"), "application/pdf")
    _clean_seed_rows(postgres_engine, data)

    try:
        with pytest.raises(
            M2SeedDataMismatchError,
            match="Existing M2 Storage object differs from seed",
        ):
            seed_m2_files(
                postgres_settings,
                manifest_path=tmp_path / "m2_manifest.json",
                m1_manifest_path=tmp_path / "m1_manifest.json",
                storage=storage,
            )
        assert not (tmp_path / "m2_manifest.json").exists()
        tenant_id = m1_deterministic_id(
            "m1-v1",
            "tenant",
            str(data["tenant_name"]),
        )
        with Session(postgres_engine) as session:
            assert (
                session.scalar(
                    select(func.count())
                    .select_from(Document)
                    .where(Document.tenant_id == tenant_id)
                )
                == 0
            )
    finally:
        _clean_seed_rows(postgres_engine, data)
