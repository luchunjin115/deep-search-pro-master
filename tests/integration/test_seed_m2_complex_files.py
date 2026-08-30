from __future__ import annotations

import hashlib
import io
from collections.abc import Generator
from pathlib import Path
from zipfile import ZipFile

import pymupdf
import pytest
from alembic import command
from alembic.config import Config
from docx import Document as WordDocument
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import Engine, delete, func, select
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.db.session import create_database_runtime
from app.models.knowledge import Document, DocumentAcl, DocumentVersion, StoredFile
from app.services.documents.parsers import DocxParser, PdfParser, XlsxParser
from app.services.documents.parsers.base import SourceLocator
from app.services.storage import LocalStorageBackend
from scripts.seed_m2_complex_files import (
    DEFAULT_COMPLEX_SEED_PATH,
    M2ComplexSeedDataMismatchError,
    build_complex_manifest,
    generate_complex_sources,
    load_complex_seed_definition,
    seed_m2_complex_files,
)
from scripts.seed_m2_files import deterministic_id

_PROJECT_ROOT = Path(__file__).resolve().parents[2]
_FIXED_ZIP_TIME = (2020, 1, 1, 0, 0, 0)
_M2_V1_BASELINE_HASHES = {
    "data/seed/m2_seed.json": "d4d9eeac120f7ea456dc48a54d70e950d4e88c4b0c9d369152332344bec6f474",
    "data/seed/m2_manifest.json": "3b522c244a85af03c265e9752216d68809142644e3c3000c89a198a1c1f45ecb",
    "scripts/m2_seed_content.py": "5075a696fe1e9e679611bfff31ea947db16e9dcc1dfc1eb158372c915e1c97a2",
    "scripts/seed_m2_files.py": "3a4fbb37a0df89714f0fcc5923da329d0f75e891cb2e3b0498ea0d37f555ae14",
}


@pytest.fixture
def postgres_settings(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Settings:
    settings = Settings(  # type: ignore[call-arg]
        _env_file=".env.example",
        app_env="test",
        local_storage_root=tmp_path / "storage",
    )
    monkeypatch.setenv("DATABASE_URL", settings.database_url)
    return settings


@pytest.fixture
def postgres_engine(postgres_settings: Settings) -> Generator[Engine, None, None]:
    runtime = create_database_runtime(postgres_settings)
    yield runtime.engine
    runtime.engine.dispose()


def _complex_ids(data: dict[str, object], entity: str) -> list[object]:
    documents = data["documents"]
    assert isinstance(documents, list)
    return [
        deterministic_id(str(data["version"]), entity, str(document["key"]))
        for document in documents
    ]


def _clean_complex_rows(engine: Engine, data: dict[str, object]) -> None:
    document_ids = _complex_ids(data, "document")
    file_ids = _complex_ids(data, "file")
    with engine.begin() as connection:
        connection.execute(delete(Document).where(Document.id.in_(document_ids)))
        connection.execute(delete(StoredFile).where(StoredFile.id.in_(file_ids)))


def test_complex_sources_are_deterministic_structured_and_keep_m2_v1_immutable() -> None:
    data = load_complex_seed_definition(DEFAULT_COMPLEX_SEED_PATH)
    first = generate_complex_sources(data)
    second = generate_complex_sources(data)
    manifest = build_complex_manifest(data, first)

    assert [source.content for source in first] == [source.content for source in second]
    assert [source.sha256 for source in first] == [source.sha256 for source in second]
    assert manifest["version"] == "m2-complex-v1"
    assert manifest["seeded_row_counts"] == {
        "files": 5,
        "documents": 5,
        "document_versions": 5,
        "document_acl": 4,
    }
    assert sum(len(item["golden_facts"]) for item in manifest["documents"]) == 10
    assert all(item["expected_route"] == "docling" for item in manifest["documents"])
    assert {
        item["key"] for item in manifest["documents"] if item["requires_ocr"]
    } == {"scanned_receiving_ticket", "visual_quality_notice"}
    for source in first:
        assert source.content
        assert len(source.definition["complexity_tags"]) >= 2
        assert b"API_KEY" not in source.content
        assert b"PRIVATE KEY" not in source.content
        for fact in source.definition["golden_facts"]:
            SourceLocator.model_validate(fact["locator"])
        if source.definition["format"] in {"docx", "xlsx"}:
            with ZipFile(io.BytesIO(source.content)) as archive:
                assert all(
                    member.date_time == _FIXED_ZIP_TIME
                    for member in archive.infolist()
                )

    for relative_path, expected_hash in _M2_V1_BASELINE_HASHES.items():
        actual_hash = hashlib.sha256((_PROJECT_ROOT / relative_path).read_bytes()).hexdigest()
        assert actual_hash == expected_hash

    sources = {source.definition["key"]: source for source in first}
    scanned = PdfParser().parse(io.BytesIO(sources["scanned_receiving_ticket"].content))
    assert scanned.page_count == 2
    assert [(page.character_count, page.image_count) for page in scanned.pages] == [
        (0, 1),
        (0, 1),
    ]
    assert [warning.code for warning in scanned.warnings] == [
        "empty_page",
        "scanned_page_suspected",
        "empty_page",
        "scanned_page_suspected",
    ]

    columns = PdfParser().parse(io.BytesIO(sources["two_column_market_brief"].content))
    assert columns.page_count == 2
    assert "德国站合成销量：139件" in columns.pages[0].text
    assert "广告预算上限：每日18 EUR" in columns.pages[0].text
    assert "DE-FRA可售库存：125件" in columns.pages[1].text
    assert "补货触发线：80件" in columns.pages[1].text

    table = PdfParser().parse(io.BytesIO(sources["merged_header_cost_table"].content))
    assert table.page_count == 1
    assert "成本构成（EUR/件）" in table.pages[0].text
    assert "合成供应商C" in table.pages[0].text
    assert "20.90" in table.pages[0].text

    visual_docx = sources["visual_quality_notice"]
    native_docx = DocxParser().parse(io.BytesIO(visual_docx.content))
    native_text = "\n".join(
        block.text for block in native_docx.blocks if block.kind == "paragraph"
    )
    assert "常规正文要求" in native_text
    assert "QC-VISUAL-17" not in native_text
    assert "IMG-D17" not in native_text
    word = WordDocument(io.BytesIO(visual_docx.content))
    assert "QC-VISUAL-17" in word.sections[0].header.paragraphs[0].text
    assert "QUALITY-LEAD" in word.sections[0].footer.paragraphs[0].text
    assert len(word.inline_shapes) == 1

    workbook_source = sources["multi_region_replenishment"]
    parsed_workbook = XlsxParser().parse(io.BytesIO(workbook_source.content))
    assert parsed_workbook.formula_count == 6
    assert [(sheet.sheet_name, sheet.row_count, sheet.column_count) for sheet in parsed_workbook.sheets] == [
        ("补货测算", 10, 7),
        ("参数说明", 4, 3),
    ]
    assert parsed_workbook.sheets[0].rows[2].cells[5].value == 100
    assert parsed_workbook.sheets[0].rows[9].cells[3].formula == "=B10+C10"
    workbook = load_workbook(io.BytesIO(workbook_source.content), data_only=False)
    try:
        assert {
            str(item) for item in workbook["补货测算"].merged_cells.ranges
        } == {"A1:A2", "B1:D1", "E1:G1", "A7:G7"}
        assert workbook["补货测算"]["G3"].value == "=E3+F3"
        assert workbook["参数说明"].freeze_panes == "A3"
    finally:
        workbook.close()

    for key in (
        "scanned_receiving_ticket",
        "two_column_market_brief",
        "merged_header_cost_table",
    ):
        pdf = pymupdf.open(stream=sources[key].content, filetype="pdf")
        try:
            assert pdf.page_count == manifest["documents"][[
                item["key"] for item in manifest["documents"]
            ].index(key)]["structure"]["page_count"]
            assert all(page.get_pixmap(matrix=pymupdf.Matrix(1, 1)).width > 0 for page in pdf)
        finally:
            pdf.close()


def test_complex_seed_is_repeatable_in_storage_and_postgresql(
    postgres_settings: Settings,
    postgres_engine: Engine,
    tmp_path: Path,
) -> None:
    command.upgrade(Config("alembic.ini"), "head")
    data = load_complex_seed_definition()
    manifest_path = tmp_path / "m2_complex_manifest.json"
    m1_manifest_path = tmp_path / "m1_manifest.json"
    _clean_complex_rows(postgres_engine, data)
    try:
        first = seed_m2_complex_files(
            postgres_settings,
            manifest_path=manifest_path,
            m1_manifest_path=m1_manifest_path,
        )
        first_bytes = manifest_path.read_bytes()
        second = seed_m2_complex_files(
            postgres_settings,
            manifest_path=manifest_path,
            m1_manifest_path=m1_manifest_path,
        )
        assert first.manifest == second.manifest
        assert first_bytes == manifest_path.read_bytes()

        file_ids = _complex_ids(data, "file")
        document_ids = _complex_ids(data, "document")
        version_ids = _complex_ids(data, "document_version")
        with Session(postgres_engine) as session:
            assert session.scalar(
                select(func.count()).select_from(StoredFile).where(StoredFile.id.in_(file_ids))
            ) == 5
            assert session.scalar(
                select(func.count()).select_from(Document).where(Document.id.in_(document_ids))
            ) == 5
            assert session.scalar(
                select(func.count()).select_from(DocumentVersion).where(
                    DocumentVersion.id.in_(version_ids)
                )
            ) == 5
            assert session.scalar(
                select(func.count()).select_from(DocumentAcl).where(
                    DocumentAcl.document_id.in_(document_ids)
                )
            ) == 4

        storage = LocalStorageBackend(postgres_settings.local_storage_root)
        for source in generate_complex_sources(data):
            with storage.open(source.storage_key) as stream:
                assert stream.read() == source.content
    finally:
        _clean_complex_rows(postgres_engine, data)


def test_complex_seed_rejects_a_conflicting_storage_object(
    postgres_settings: Settings,
    postgres_engine: Engine,
    tmp_path: Path,
) -> None:
    command.upgrade(Config("alembic.ini"), "head")
    data = load_complex_seed_definition()
    source = generate_complex_sources(data)[0]
    storage = LocalStorageBackend(postgres_settings.local_storage_root)
    storage.put(source.storage_key, io.BytesIO(b"conflicting bytes"), "application/pdf")
    _clean_complex_rows(postgres_engine, data)
    try:
        with pytest.raises(
            M2ComplexSeedDataMismatchError,
            match="Existing M2 Storage object differs from seed",
        ):
            seed_m2_complex_files(
                postgres_settings,
                manifest_path=tmp_path / "m2_complex_manifest.json",
                m1_manifest_path=tmp_path / "m1_manifest.json",
                storage=storage,
            )
        assert not (tmp_path / "m2_complex_manifest.json").exists()
    finally:
        _clean_complex_rows(postgres_engine, data)
