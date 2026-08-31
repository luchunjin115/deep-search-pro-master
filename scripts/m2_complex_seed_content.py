"""Deterministic builders for the m2-complex-v1 synthetic evaluation corpus."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

from docx import Document as WordDocument
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import (  # type: ignore[import-untyped]
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
from reportlab.lib.utils import ImageReader  # type: ignore[import-untyped]
from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
from reportlab.pdfbase.cidfonts import UnicodeCIDFont  # type: ignore[import-untyped]
from reportlab.pdfgen.canvas import Canvas  # type: ignore[import-untyped]

from scripts.m2_seed_content import _canonicalize_zip

_FIXED_TIME = datetime(2026, 8, 30, tzinfo=UTC)
_VERSION = "m2-complex-v1"


class M2ComplexSeedContentError(ValueError):
    """A reviewed complex seed definition cannot be rendered safely."""


def build_complex_source_bytes(document: dict[str, Any], disclaimer: str) -> bytes:
    """Render one complex source without using Docling or a Parser Router."""

    key = document["key"]
    if key == "scanned_receiving_ticket":
        return _build_scanned_pdf(document, disclaimer)
    if key == "two_column_market_brief":
        return _build_two_column_pdf(document, disclaimer)
    if key == "merged_header_cost_table":
        return _build_complex_table_pdf(document, disclaimer)
    if key == "visual_quality_notice":
        return _build_visual_docx(document, disclaimer)
    if key == "multi_region_replenishment":
        return _build_complex_xlsx(document, disclaimer)
    raise M2ComplexSeedContentError("Complex seed document key is not supported")


def _new_pdf(document: dict[str, Any], output: io.BytesIO) -> Canvas:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    canvas = Canvas(output, pagesize=A4, pageCompression=1, invariant=1)
    canvas.setTitle(document["title"])
    canvas.setAuthor("Deep Search Pro synthetic demo")
    canvas.setSubject(f"{_VERSION} complex evaluation corpus")
    return canvas


def _build_scanned_pdf(document: dict[str, Any], disclaimer: str) -> bytes:
    output = io.BytesIO()
    canvas = _new_pdf(document, output)
    width, height = A4
    pages = document["content"]["pages"]
    if len(pages) != 2:
        raise M2ComplexSeedContentError("Scanned PDF requires exactly two pages")

    for page_number, page in enumerate(pages, start=1):
        image = Image.new("RGB", (1240, 1754), "#F7F4EC")
        drawing = ImageDraw.Draw(image)
        title_font = ImageFont.load_default(size=40)
        body_font = ImageFont.load_default(size=30)
        small_font = ImageFont.load_default(size=20)
        drawing.rectangle((70, 70, 1170, 1684), outline="#222222", width=5)
        drawing.text(
            (115, 120), page["image_lines"][0], fill="#111111", font=title_font
        )
        y = 280
        for line in page["image_lines"][1:]:
            drawing.text((130, y), line, fill="#111111", font=body_font)
            drawing.line((120, y + 55, 1110, y + 55), fill="#B7B0A4", width=2)
            y += 165
        drawing.text(
            (115, 1560),
            f"{_VERSION} | PAGE {page_number} | SYNTHETIC DEMO DATA",
            fill="#555555",
            font=small_font,
        )
        image_bytes = io.BytesIO()
        image.save(image_bytes, format="PNG", optimize=False, compress_level=9)
        image_bytes.seek(0)
        canvas.drawImage(
            ImageReader(image_bytes),
            36,
            42,
            width=width - 72,
            height=height - 84,
            preserveAspectRatio=True,
            anchor="c",
        )
        canvas.showPage()
    canvas.save()
    return output.getvalue()


def _build_two_column_pdf(document: dict[str, Any], disclaimer: str) -> bytes:
    output = io.BytesIO()
    canvas = _new_pdf(document, output)
    width, height = A4
    for page_number, page in enumerate(document["content"]["pages"], start=1):
        canvas.setFont("STSong-Light", 9)
        canvas.drawString(48, height - 34, "LUMORIVA | 合成市场简报 | 重复页眉")
        canvas.drawRightString(width - 48, height - 34, f"{_VERSION} | {page_number}/2")
        canvas.setStrokeColorRGB(0.78, 0.78, 0.78)
        canvas.line(48, height - 44, width - 48, height - 44)
        canvas.setFont("STSong-Light", 20)
        canvas.setFillColorRGB(0.04, 0.15, 0.27)
        canvas.drawString(48, height - 82, document["title"])
        canvas.setFillColorRGB(0, 0, 0)

        gutter = 24
        column_width = (width - 96 - gutter) / 2
        left_x = 48
        right_x = left_x + column_width + gutter
        canvas.setStrokeColorRGB(0.88, 0.88, 0.88)
        canvas.line(width / 2, height - 112, width / 2, 90)
        for x, heading, lines in (
            (left_x, page["left_heading"], page["left_lines"]),
            (right_x, page["right_heading"], page["right_lines"]),
        ):
            canvas.setFont("STSong-Light", 14)
            canvas.setFillColorRGB(0.18, 0.45, 0.71)
            canvas.drawString(x, height - 130, heading)
            canvas.setFillColorRGB(0, 0, 0)
            y = height - 174
            canvas.setFont("STSong-Light", 11)
            for line in lines:
                canvas.drawString(x + 8, y, f"• {line}")
                y -= 42
        canvas.setFont("STSong-Light", 8)
        canvas.setFillColorRGB(0.35, 0.35, 0.35)
        canvas.drawString(48, 48, f"{disclaimer} | 重复页脚")
        canvas.showPage()
    canvas.save()
    return output.getvalue()


def _build_complex_table_pdf(document: dict[str, Any], disclaimer: str) -> bytes:
    output = io.BytesIO()
    canvas = _new_pdf(document, output)
    _width, height = A4
    content = document["content"]
    x0, y0 = 48, height - 180
    widths = [132, 88, 88, 88, 105]
    row_height = 34
    total_width = sum(widths)

    canvas.setFont("STSong-Light", 20)
    canvas.drawString(48, height - 68, document["title"])
    canvas.setFont("STSong-Light", 9)
    canvas.drawString(48, height - 92, f"{_VERSION} | 合成演示数据 | 无边框业务表")

    canvas.setFillColorRGB(0.12, 0.31, 0.47)
    canvas.rect(x0, y0, widths[0], row_height, fill=1, stroke=0)
    canvas.rect(
        x0 + widths[0], y0, total_width - widths[0], row_height, fill=1, stroke=0
    )
    canvas.setFillColorRGB(1, 1, 1)
    canvas.setFont("STSong-Light", 10)
    canvas.drawCentredString(x0 + widths[0] / 2, y0 + 11, content["groups"][0]["label"])
    canvas.drawCentredString(
        x0 + widths[0] + (total_width - widths[0]) / 2,
        y0 + 11,
        content["groups"][1]["label"],
    )

    y = y0 - row_height
    canvas.setFillColorRGB(0.85, 0.91, 0.95)
    canvas.rect(x0, y, total_width, row_height, fill=1, stroke=0)
    canvas.setFillColorRGB(0.08, 0.14, 0.2)
    cursor = x0
    for label, cell_width in zip(content["columns"], widths):
        canvas.drawCentredString(cursor + cell_width / 2, y + 11, label)
        cursor += cell_width

    canvas.setFont("STSong-Light", 10)
    for row_index, row in enumerate(content["rows"], start=1):
        y -= row_height
        if row_index % 2 == 0:
            canvas.setFillColorRGB(0.96, 0.97, 0.98)
            canvas.rect(x0, y, total_width, row_height, fill=1, stroke=0)
        canvas.setFillColorRGB(0.08, 0.08, 0.08)
        cursor = x0
        for value, cell_width in zip(row, widths):
            canvas.drawCentredString(cursor + cell_width / 2, y + 11, value)
            cursor += cell_width

    canvas.setFont("STSong-Light", 8)
    canvas.setFillColorRGB(0.35, 0.35, 0.35)
    canvas.drawString(48, 48, disclaimer)
    canvas.showPage()
    canvas.save()
    return output.getvalue()


def _build_visual_docx(document: dict[str, Any], disclaimer: str) -> bytes:
    word = WordDocument()
    _apply_docx_preset(word)
    properties = word.core_properties
    properties.title = document["title"]
    properties.subject = f"{_VERSION} complex evaluation corpus"
    properties.author = "Deep Search Pro synthetic demo"
    properties.keywords = "synthetic,demo,complex,ocr"
    properties.comments = disclaimer
    properties.created = _FIXED_TIME
    properties.modified = _FIXED_TIME
    properties.last_printed = _FIXED_TIME
    properties.revision = 1

    content = document["content"]
    header = word.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_run(header, content["header_text"], size=8, color="666666")
    footer = word.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(footer, content["footer_text"], size=8, color="666666")

    title = word.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _add_run(title, document["title"], size=22, color="0B2545", bold=True)
    subtitle = word.add_paragraph()
    _add_run(subtitle, f"{_VERSION} | 合成复杂文档评估资料", size=10, color="555555")
    notice = word.add_paragraph(disclaimer)
    for run in notice.runs:
        _set_run_font(run, size=9, color="7A3E00")

    for section in content["sections"]:
        word.add_heading(section["heading"], level=1)
        for text in section["paragraphs"]:
            paragraph = word.add_paragraph(text)
            for run in paragraph.runs:
                _set_run_font(run, size=11)

    word.add_heading("3. 图片质检卡", level=1)
    image_bytes = _make_text_card(content["image_lines"])
    picture = word.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(image_bytes, width=Inches(5.8))
    caption = word.add_paragraph("图1：文字只存在于图片中的合成质检卡")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        _set_run_font(run, size=9, color="666666")

    output = io.BytesIO()
    word.save(output)
    return _canonicalize_zip(output.getvalue())


def _make_text_card(lines: list[str]) -> io.BytesIO:
    image = Image.new("RGB", (1200, 520), "#FFF7E8")
    drawing = ImageDraw.Draw(image)
    drawing.rounded_rectangle(
        (25, 25, 1175, 495), radius=24, outline="#B45309", width=6
    )
    title_font = ImageFont.load_default(size=36)
    body_font = ImageFont.load_default(size=30)
    for index, line in enumerate(lines):
        drawing.text(
            (75, 75 + index * 115),
            line,
            fill="#3A2410",
            font=title_font if index == 0 else body_font,
        )
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=False, compress_level=9)
    output.seek(0)
    return output


def _apply_docx_preset(word: Any) -> None:
    section = word.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = word.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = word.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def _add_run(
    paragraph: Any,
    text: str,
    *,
    size: int,
    color: str | None = None,
    bold: bool | None = None,
) -> None:
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, color=color, bold=bold)


def _set_run_font(
    run: Any,
    *,
    size: int,
    color: str | None = None,
    bold: bool | None = None,
) -> None:
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _build_complex_xlsx(document: dict[str, Any], disclaimer: str) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = document["title"]
    workbook.properties.subject = f"{_VERSION} complex evaluation corpus"
    workbook.properties.creator = "Deep Search Pro synthetic demo"
    workbook.properties.lastModifiedBy = "Deep Search Pro synthetic demo"
    workbook.properties.description = disclaimer
    workbook.properties.created = _FIXED_TIME.replace(tzinfo=None)
    workbook.properties.modified = _FIXED_TIME.replace(tzinfo=None)

    for definition in document["content"]["sheets"]:
        sheet = workbook.create_sheet(definition["name"])
        sheet.sheet_view.showGridLines = False
        for row in definition["rows"]:
            sheet.append(row)
        for merged_range in definition["merged_ranges"]:
            sheet.merge_cells(merged_range)
        _style_complex_sheet(sheet)

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return _canonicalize_zip(output.getvalue())


def _style_complex_sheet(sheet: Any) -> None:
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    medium_fill = PatternFill("solid", fgColor="D9EAF7")
    section_fill = PatternFill("solid", fgColor="F4B183")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    dark_bold = Font(name="Calibri", size=10, bold=True, color="1F2937")
    body_font = Font(name="Calibri", size=10, color="222222")
    thin = Side(style="thin", color="D1D5DB")

    for row in sheet.iter_rows():
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.value is not None:
                cell.border = Border(bottom=thin)
    if sheet.title == "补货测算":
        sheet.freeze_panes = "B3"
        for cell in sheet[1]:
            cell.fill = dark_fill
            cell.font = white_bold
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for cell in sheet[2]:
            cell.fill = medium_fill
            cell.font = dark_bold
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet["A7"].fill = section_fill
        sheet["A7"].font = dark_bold
        for cell in sheet[8]:
            cell.fill = medium_fill
            cell.font = dark_bold
        widths: tuple[int, ...] = (22, 14, 14, 14, 14, 14, 16)
    else:
        sheet.freeze_panes = "A3"
        sheet["A1"].fill = dark_fill
        sheet["A1"].font = white_bold
        for cell in sheet[2]:
            cell.fill = medium_fill
            cell.font = dark_bold
        widths = (20, 16, 32)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row_number in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = 24
