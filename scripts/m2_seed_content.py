"""Deterministic binary builders for the versioned M2 synthetic corpus."""

from __future__ import annotations

import csv
import io
import re
from datetime import UTC, datetime
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document as WordDocument
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
from reportlab.pdfbase.cidfonts import UnicodeCIDFont  # type: ignore[import-untyped]
from reportlab.pdfgen.canvas import Canvas  # type: ignore[import-untyped]

_FIXED_TIME = datetime(2026, 8, 29, tzinfo=UTC)
_ZIP_TIMESTAMP = (2020, 1, 1, 0, 0, 0)
_FIXED_W3CDTF = b"2026-08-29T00:00:00Z"
_DOCX_TABLE_WIDTH_DXA = 9360
_DOCX_TABLE_INDENT_DXA = 120
_DISCLAIMER = "全部内容均为合成演示数据。"


class M2SeedContentError(ValueError):
    """A trusted M2 seed definition cannot be rendered safely."""


def build_source_bytes(document: dict[str, Any], disclaimer: str) -> bytes:
    """Build one supported source artifact from its validated seed definition."""

    source_format = document["format"]
    if source_format == "pdf":
        return _build_pdf(document, disclaimer)
    if source_format == "docx":
        return _build_docx(document, disclaimer)
    if source_format == "xlsx":
        return _build_xlsx(document, disclaimer)
    if source_format == "csv":
        return _build_csv(document, disclaimer)
    raise M2SeedContentError("M2 seed source format is not supported")


def _build_pdf(document: dict[str, Any], disclaimer: str) -> bytes:
    content = document["content"]
    pages = content.get("pages")
    if not isinstance(pages, list) or not pages:
        raise M2SeedContentError("PDF seed requires at least one page")

    output = io.BytesIO()
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    canvas = Canvas(
        output,
        pagesize=A4,
        pageCompression=1,
        invariant=1,
    )
    canvas.setTitle(document["title"])
    canvas.setAuthor("Deep Search Pro synthetic demo")
    canvas.setSubject("m2-v1 synthetic knowledge corpus")
    width, height = A4

    for page_number, page in enumerate(pages, start=1):
        canvas.setFont("STSong-Light", 20)
        canvas.drawString(54, height - 66, page["title"])
        canvas.setFont("STSong-Light", 9)
        canvas.drawRightString(width - 54, height - 40, f"m2-v1 | 第{page_number}页")
        y = height - 106
        for section in page["sections"]:
            canvas.setFont("STSong-Light", 14)
            canvas.drawString(54, y, section["heading"])
            y -= 28
            canvas.setFont("STSong-Light", 11)
            for line in section["lines"]:
                if y < 78:
                    raise M2SeedContentError("PDF seed page content exceeds fixed layout")
                canvas.drawString(66, y, line)
                y -= 21
            y -= 10
        canvas.setFont("STSong-Light", 8)
        canvas.setFillColorRGB(0.35, 0.35, 0.35)
        canvas.drawString(54, 40, f"{disclaimer} | {_DISCLAIMER}")
        canvas.setFillColorRGB(0, 0, 0)
        canvas.showPage()
    canvas.save()
    return output.getvalue()


def _build_docx(document: dict[str, Any], disclaimer: str) -> bytes:
    word = WordDocument()
    _apply_docx_preset(word)
    _set_docx_metadata(word, document)
    _add_docx_header_footer(word, document["title"])
    _add_docx_title_block(word, document, disclaimer)

    for section in document["content"]["sections"]:
        level = int(section["level"])
        word.add_heading(section["heading"], level=level)
        for paragraph_text in section.get("paragraphs", []):
            paragraph = word.add_paragraph(paragraph_text)
            _set_all_run_fonts(paragraph)
        table_definition = section.get("table")
        if table_definition is not None:
            _add_docx_table(
                word,
                headers=table_definition["headers"],
                rows=table_definition["rows"],
                widths=tuple(table_definition["widths_dxa"]),
            )

    output = io.BytesIO()
    word.save(output)
    return _canonicalize_zip(output.getvalue())


def _apply_docx_preset(document: DocxDocument) -> None:
    """Apply the compact_reference_guide preset as explicit Word tokens."""

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    _set_style_font(normal, "Calibri", 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 14, 7),
        "Heading 3": (12, "1F4D78", 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[style_name]
        _set_style_font(style, "Calibri", size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25


def _set_style_font(style: Any, latin_name: str, size: int) -> None:
    style.font.name = latin_name
    style.font.size = Pt(size)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin_name)
    fonts.set(qn("w:hAnsi"), latin_name)
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _set_docx_metadata(document: DocxDocument, definition: dict[str, Any]) -> None:
    properties = document.core_properties
    properties.title = definition["title"]
    properties.subject = "m2-v1 synthetic knowledge corpus"
    properties.author = "Deep Search Pro synthetic demo"
    properties.keywords = "synthetic,demo,m2-v1"
    properties.comments = _DISCLAIMER
    properties.created = _FIXED_TIME
    properties.modified = _FIXED_TIME
    properties.last_printed = _FIXED_TIME
    properties.revision = 1


def _add_docx_header_footer(document: DocxDocument, title: str) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run(f"Deep Search Pro | {title}")
    _set_run_font(run, size=8, color="666666")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    run = footer.add_run("m2-v1 | 合成演示资料")
    _set_run_font(run, size=8, color="666666")


def _add_docx_title_block(
    document: DocxDocument,
    definition: dict[str, Any],
    disclaimer: str,
) -> None:
    """Use a restrained memo_masthead without a decorative bottom rule."""

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(definition["title"])
    _set_run_font(run, size=22, color="0B2545", bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("m2-v1 合成知识文件")
    _set_run_font(run, size=11, color="555555")

    for label, value in (
        ("版本", "m2-v1"),
        ("资料性质", _DISCLAIMER),
        ("重要说明", disclaimer),
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}：")
        _set_run_font(label_run, size=10, bold=True)
        value_run = paragraph.add_run(value)
        _set_run_font(value_run, size=10)


def _add_docx_table(
    document: DocxDocument,
    *,
    headers: list[str],
    rows: list[list[str]],
    widths: tuple[int, ...],
) -> None:
    if len(headers) != len(widths) or sum(widths) != _DOCX_TABLE_WIDTH_DXA:
        raise M2SeedContentError("DOCX seed table geometry is invalid")
    if any(len(row) != len(headers) for row in rows):
        raise M2SeedContentError("DOCX seed table rows do not match headers")

    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_docx_table_geometry(table, widths)

    values = [headers, *rows]
    for row_number, row_values in enumerate(values):
        for column_number, value in enumerate(row_values):
            cell = table.cell(row_number, column_number)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if column_number == 0
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.25
            _set_all_run_fonts(paragraph)
            if row_number == 0:
                for run in paragraph.runs:
                    run.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EEF5")
                cell._tc.get_or_add_tcPr().append(shading)


def _set_docx_table_geometry(table: Any, widths: tuple[int, ...]) -> None:
    table_xml = table._tbl
    properties = table_xml.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.insert(0, table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(_DOCX_TABLE_WIDTH_DXA))

    table_indent = properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        properties.append(table_indent)
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), str(_DOCX_TABLE_INDENT_DXA))

    for grid_column, width in zip(table_xml.tblGrid.gridCol_lst, widths):
        grid_column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.get_or_add_tcW()
            cell_width.type = "dxa"
            cell_width.w = width
            margins = cell_properties.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell_properties.append(margins)
            for side, value in (
                ("top", 80),
                ("bottom", 80),
                ("start", 120),
                ("end", 120),
            ):
                element = margins.find(qn(f"w:{side}"))
                if element is None:
                    element = OxmlElement(f"w:{side}")
                    margins.append(element)
                element.set(qn("w:w"), str(value))
                element.set(qn("w:type"), "dxa")


def _set_all_run_fonts(paragraph: Any) -> None:
    for run in paragraph.runs:
        _set_run_font(run, size=11)


def _set_run_font(
    run: Any,
    *,
    size: int,
    color: str | None = None,
    bold: bool | None = None,
) -> None:
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "Microsoft YaHei"
    )
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _build_xlsx(document: dict[str, Any], disclaimer: str) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = document["title"]
    workbook.properties.subject = "m2-v1 synthetic knowledge corpus"
    workbook.properties.creator = "Deep Search Pro synthetic demo"
    workbook.properties.lastModifiedBy = "Deep Search Pro synthetic demo"
    workbook.properties.description = f"{_DISCLAIMER} {disclaimer}"
    workbook.properties.created = _FIXED_TIME.replace(tzinfo=None)
    workbook.properties.modified = _FIXED_TIME.replace(tzinfo=None)

    for sheet_definition in document["content"]["sheets"]:
        sheet = workbook.create_sheet(sheet_definition["name"])
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        headers = sheet_definition["headers"]
        sheet.append(headers)
        for row in sheet_definition["rows"]:
            sheet.append(row)
        for formula in sheet_definition.get("formulas", []):
            sheet[formula["coordinate"]] = formula["formula"]
        _style_xlsx_sheet(sheet, sheet_definition["widths"])
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return _canonicalize_zip(output.getvalue())


def _style_xlsx_sheet(sheet: Any, widths: list[int]) -> None:
    if len(widths) != sheet.max_column:
        raise M2SeedContentError("XLSX seed column widths do not match data")
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=10, color="222222")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = 22


def _build_csv(document: dict[str, Any], disclaimer: str) -> bytes:
    content = document["content"]
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(content["headers"])
    writer.writerows(content["rows"])
    writer.writerow([])
    writer.writerow(["说明", _DISCLAIMER, disclaimer])
    return output.getvalue().encode("utf-8-sig")


def _canonicalize_zip(source: bytes) -> bytes:
    """Remove volatile ZIP and core-property times from OOXML packages."""

    output = io.BytesIO()
    with (
        ZipFile(io.BytesIO(source), "r") as input_archive,
        ZipFile(output, "w", compression=ZIP_DEFLATED, compresslevel=9) as archive,
    ):
        for member in sorted(input_archive.infolist(), key=lambda item: item.filename):
            info = ZipInfo(member.filename, date_time=_ZIP_TIMESTAMP)
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0o600 << 16
            content = input_archive.read(member.filename)
            if member.filename == "docProps/core.xml":
                for property_name in (b"created", b"modified"):
                    content = re.sub(
                        rb"(<dcterms:"
                        + property_name
                        + rb"\b[^>]*>)[^<]*(</dcterms:"
                        + property_name
                        + rb">)",
                        rb"\g<1>" + _FIXED_W3CDTF + rb"\g<2>",
                        content,
                    )
            archive.writestr(info, content)
    return output.getvalue()
